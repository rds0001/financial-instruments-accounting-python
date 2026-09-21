# Copyright 2026 RiskDataScience GmbH
# SPDX-License-Identifier: Apache-2.0
import json
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .contracts import CONTRACT_VERSION, WORKBOOKS
from .excel_io import read_table, sha256, table_columns


class ValidationRejected(ValueError):
    def __init__(self, issues: list[dict]):
        self.issues = issues
        super().__init__(f"Datensatz abgelehnt: {len(issues)} Fehler")


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _as_datetime(value) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def select_official_snapshot(
    tables: dict[str, list[dict]], reporting_date: str, knowledge_time: str
) -> tuple[dict[str, list[dict]], list[dict]]:
    """Select the one official, active record valid and known at the run cut-offs."""
    selected: dict[str, list[dict]] = {}
    issues: list[dict] = []
    cut_off_date, cut_off_time = _as_date(reporting_date), _as_datetime(knowledge_time)
    for table_name, rows in tables.items():
        by_key: dict[object, list[dict]] = defaultdict(list)
        for row in rows:
            by_key[row.get("business_key")].append(row)
        selected[table_name] = []
        for business_key, versions in by_key.items():
            applicable = []
            for row in versions:
                try:
                    valid = (
                        _as_date(row.get("valid_from"))
                        <= cut_off_date
                        <= _as_date(row.get("valid_to"))
                    )
                    known = (
                        _as_datetime(row.get("known_from"))
                        <= cut_off_time
                        <= _as_datetime(row.get("known_to"))
                    )
                except (TypeError, ValueError, OverflowError):
                    continue
                if (
                    row.get("is_official") is True
                    and row.get("record_status") == "ACTIVE"
                    and valid
                    and known
                ):
                    applicable.append(row)
            if len(applicable) != 1:
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "OFFICIAL_SNAPSHOT_CARDINALITY",
                        "object": f"{table_name}:{business_key}",
                        "message": f"Zum Stichtag/Kenntniszeitpunkt wurden {len(applicable)} statt genau einer Official-Version gefunden",
                    }
                )
            else:
                selected[table_name].append(applicable[0])
    return selected, issues


def load_and_validate(
    dataset: Path, reporting_date: str | None = None, knowledge_time: str | None = None
) -> tuple[dict[str, list[dict]], list[dict]]:
    issues: list[dict] = []
    tables: dict[str, list[dict]] = {}
    for filename, (table_name, columns) in WORKBOOKS.items():
        path = dataset / filename
        if not path.is_file():
            issues.append(
                {
                    "severity": "FATAL",
                    "code": "MISSING_WORKBOOK",
                    "object": filename,
                    "message": "Pflichtdatei fehlt",
                }
            )
            continue
        try:
            rows = read_table(path, table_name)
            tables[table_name] = rows
            expected = table_columns(columns)
            wb = load_workbook(path, read_only=False, data_only=False)
            ws = wb[table_name.upper()]
            actual = [cell.value for cell in ws[1]]
            if actual != expected:
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "SCHEMA_COLUMNS",
                        "object": filename,
                        "message": f"expected={expected}; actual={actual}",
                    }
                )
            if f"tbl_in_{table_name}" not in ws.tables:
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "MISSING_EXCEL_TABLE",
                        "object": filename,
                        "message": "Strukturierte Eingabetabelle fehlt",
                    }
                )
            ids = [r.get("record_id") for r in rows]
            if None in ids or len(ids) != len(set(ids)):
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "PRIMARY_KEY",
                        "object": filename,
                        "message": "record_id fehlt oder ist doppelt",
                    }
                )
            for row in rows:
                if any(
                    row.get(field) in (None, "")
                    for field in (
                        "record_id",
                        "business_key",
                        "version_no",
                        "as_of_date",
                        "valid_from",
                        "valid_to",
                        "known_from",
                        "known_to",
                        "record_status",
                    )
                ):
                    issues.append(
                        {
                            "severity": "ERROR",
                            "code": "MANDATORY_METADATA",
                            "object": row.get("record_id") or filename,
                            "message": "Bitemporale Pflichtmetadaten fehlen",
                        }
                    )
                if str(row.get("valid_from")) > str(row.get("valid_to")) or str(
                    row.get("known_from")
                ) > str(row.get("known_to")):
                    issues.append(
                        {
                            "severity": "ERROR",
                            "code": "INVALID_TIME_INTERVAL",
                            "object": row.get("record_id"),
                            "message": "Ungültiges Gültigkeits-/Kenntnisintervall",
                        }
                    )
        except Exception as exc:
            issues.append(
                {
                    "severity": "FATAL",
                    "code": "WORKBOOK_READ",
                    "object": filename,
                    "message": str(exc),
                }
            )
    manifest_path = dataset / "profile_manifest.json"
    if manifest_path.is_file():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("contract_version") != CONTRACT_VERSION:
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "CONTRACT_VERSION",
                    "object": "profile_manifest.json",
                    "message": "Vertragsversion weicht ab",
                }
            )
        if set(manifest.get("files", {})) != set(WORKBOOKS):
            raise ValidationRejected(
                [
                    {
                        "severity": "FATAL",
                        "code": "PROFILE_FILE_SET",
                        "object": "profile_manifest.json",
                        "message": "Seal must cover exactly the 19 canonical workbooks",
                    }
                ]
            )
        for name, meta in manifest.get("files", {}).items():
            path = dataset / name
            if not path.is_file() or sha256(path) != meta["sha256"]:
                issues.append(
                    {
                        "severity": "FATAL",
                        "code": "PROFILE_HASH",
                        "object": name,
                        "message": "Versiegelung verletzt",
                    }
                )
        declared_counts = manifest.get("table_row_counts", {})
        actual_counts = {name: len(rows) for name, rows in tables.items()}
        if declared_counts and declared_counts != actual_counts:
            issues.append(
                {
                    "severity": "FATAL",
                    "code": "PROFILE_ROW_COUNTS",
                    "object": "profile_manifest.json",
                    "message": f"declared={declared_counts}; actual={actual_counts}",
                }
            )
        expected_files = set(WORKBOOKS) | {"profile_manifest.json"}
        actual_files = {p.name for p in dataset.iterdir() if p.is_file()}
        if actual_files != expected_files:
            issues.append(
                {
                    "severity": "FATAL",
                    "code": "PROFILE_FILE_SET",
                    "object": "profile_manifest.json",
                    "message": f"unexpected={sorted(actual_files - expected_files)} missing={sorted(expected_files - actual_files)}",
                }
            )
    return validate_tables(tables, reporting_date, knowledge_time, issues)


def validate_tables(tables, reporting_date=None, knowledge_time=None, issues=None):
    issues = list(issues or [])
    run_control = tables.get("run_control", [])
    if reporting_date is None and run_control:
        reporting_date = str(run_control[0].get("reporting_date"))
    if knowledge_time is None and run_control:
        knowledge_time = str(run_control[0].get("knowledge_time"))
    if reporting_date and knowledge_time:
        tables, snapshot_issues = select_official_snapshot(tables, reporting_date, knowledge_time)
        issues.extend(snapshot_issues)
    else:
        issues.append(
            {
                "severity": "FATAL",
                "code": "MISSING_RUN_CUTOFF",
                "object": "run_control",
                "message": "Stichtag oder Kenntniszeitpunkt fehlt",
            }
        )
    if (
        tables.get("run_control")
        and tables["run_control"][0].get("contract_version") != CONTRACT_VERSION
    ):
        issues.append(
            {
                "severity": "ERROR",
                "code": "CONTRACT_VERSION",
                "object": "run_control",
                "message": "Nicht unterstützte Vertragsversion",
            }
        )
    scenarios = tables.get("scenarios", [])
    if scenarios and abs(sum(float(r["probability"]) for r in scenarios) - 1.0) > 1e-9:
        issues.append(
            {
                "severity": "ERROR",
                "code": "SCENARIO_WEIGHT_SUM",
                "object": "scenarios",
                "message": "Wahrscheinlichkeiten summieren sich nicht zu 1",
            }
        )
    instrument_ids = {r.get("instrument_id") for r in tables.get("instruments", [])}
    for table in (
        "classification_inputs",
        "opening_balances",
        "cashflows",
        "collateral",
        "staging_inputs",
        "pd_curves",
        "lgd_profiles",
        "ead_profiles",
        "events",
        "market_data",
    ):
        for row in tables.get(table, []):
            if row.get("instrument_id") not in instrument_ids:
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "BROKEN_REFERENCE",
                        "object": table,
                        "message": str(row.get("instrument_id")),
                    }
                )
    for table in ("classification_inputs", "opening_balances", "staging_inputs", "market_data"):
        counts = Counter(row.get("instrument_id") for row in tables.get(table, []))
        if set(counts) != instrument_ids or any(value != 1 for value in counts.values()):
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "INSTRUMENT_CARDINALITY",
                    "object": table,
                    "message": "Jedes Instrument muss genau einen Datensatz besitzen",
                }
            )
    for row in tables.get("pd_curves", []):
        if not 0 <= float(row["marginal_pd"]) <= 1:
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "PD_RANGE",
                    "object": row["record_id"],
                    "message": "PD außerhalb [0,1]",
                }
            )
    for row in tables.get("lgd_profiles", []):
        if not 0 <= float(row["lgd"]) <= 1:
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "LGD_RANGE",
                    "object": row["record_id"],
                    "message": "LGD außerhalb [0,1]",
                }
            )
    for row in tables.get("ead_profiles", []):
        if (
            float(row["ead"]) < 0
            or not 0 <= float(row["ccf"]) <= 1
            or not 0 <= float(row["prepayment_rate"]) <= 1
        ):
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "EAD_CCF_RANGE",
                    "object": row["record_id"],
                    "message": "EAD/CCF/Prepayment außerhalb zulässiger Werte",
                }
            )
    pd_groups = defaultdict(float)
    for row in tables.get("pd_curves", []):
        pd_groups[(row["instrument_id"], row["scenario_id"])] += float(row["marginal_pd"])
    for key, total in pd_groups.items():
        if total > 1 + 1e-9:
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "PD_SURVIVAL",
                    "object": str(key),
                    "message": "Summe marginaler PD übersteigt 1",
                }
            )
    scenario_ids = {row.get("scenario_id") for row in scenarios}
    for table in ("pd_curves", "lgd_profiles", "ead_profiles"):
        for row in tables.get(table, []):
            if row.get("scenario_id") not in scenario_ids:
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "SCENARIO_REFERENCE",
                        "object": row.get("record_id"),
                        "message": str(row.get("scenario_id")),
                    }
                )
            if (
                row.get("validation_status") not in {"VALIDATED", "CONDITIONALLY_VALIDATED"}
                or not row.get("approval_id")
                or not row.get("model_owner")
            ):
                issues.append(
                    {
                        "severity": "ERROR",
                        "code": "MODEL_GOVERNANCE",
                        "object": row.get("record_id"),
                        "message": "Modell nicht verwendbar/freigegeben",
                    }
                )
    for row in tables.get("backtesting", []):
        if row.get("model_type") not in {
            "PD",
            "LGD",
            "EAD",
            "CCF",
            "STAGING",
            "SCENARIO",
            "OVERLAY",
        }:
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "BACKTEST_TYPE",
                    "object": row.get("record_id"),
                    "message": str(row.get("model_type")),
                }
            )
        if float(row.get("exposure_weight") or 0) <= 0:
            issues.append(
                {
                    "severity": "ERROR",
                    "code": "BACKTEST_WEIGHT",
                    "object": row.get("record_id"),
                    "message": "Gewicht muss positiv sein",
                }
            )
    mapping_keys = [
        (
            row.get("mapping_type"),
            row.get("external_value"),
            row.get("valid_from"),
            row.get("valid_to"),
        )
        for row in tables.get("mappings", [])
    ]
    if len(mapping_keys) != len(set(mapping_keys)):
        issues.append(
            {
                "severity": "ERROR",
                "code": "AMBIGUOUS_MAPPING",
                "object": "mappings",
                "message": "Mehrdeutige Mappings",
            }
        )
    fatals = [x for x in issues if x["severity"] in {"FATAL", "ERROR"}]
    if fatals:
        raise ValidationRejected(issues)
    return tables, issues
