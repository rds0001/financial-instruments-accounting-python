# Copyright 2026 RiskDataScience GmbH
# SPDX-License-Identifier: Apache-2.0
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .contracts import COMMON, CONTRACT_VERSION, DOMAINS


def _safe(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def table_columns(specific: list[str]) -> list[str]:
    return list(dict.fromkeys(COMMON + specific))


def write_contract_workbook(
    path: Path, table_name: str, specific: list[str], rows: list[dict], purpose: str
) -> None:
    wb = Workbook()
    fixed_time = datetime(2026, 9, 2, 10, 0, 0)
    wb.properties.created = fixed_time
    wb.properties.modified = fixed_time
    readme = wb.active
    readme.title = "README"
    readme.append(["IFRS-9 canonical input workbook", path.name])
    readme.append(["Purpose", purpose])
    readme.append(["Contract version", CONTRACT_VERSION])
    readme.append(
        ["Owner", "Synthetic Reference Data Owner / institution-specific replacement required"]
    )
    readme.append(
        [
            "Rules",
            "No formulas; typed values; append-only/bitemporal metadata; formula-injection protected.",
        ]
    )
    readme["A1"].font = Font(bold=True, size=14)
    dictionary = wb.create_sheet("DATA_DICTIONARY")
    dictionary.append(
        [
            "field",
            "type",
            "nullable",
            "unit",
            "key",
            "domain",
            "time_reference",
            "origin",
            "description",
        ]
    )
    columns = table_columns(specific)
    bools = {
        "liability",
        "equity",
        "held_for_trading",
        "fair_value_option",
        "fvoci_equity_election",
        "hybrid_contract",
        "embedded_derivative_separated",
        "simplified",
        "poci",
        "overlay_approved",
        "overlay_double_count_checked",
        "regular_way",
        "electronic_payment_election",
        "payment_irrevocable",
        "withdrawal_access_practical",
        "settlement_risk_insignificant",
        "settlement_short_period",
        "own_use",
        "interface_approved",
        "nature_dependent_contract",
        "purchaser_net_buyer",
        "quantity_consistent_expected_usage",
        "sale_timing_uncontrollable",
        "repurchase_within_reasonable_time",
        "regulatory_default_flag",
        "contractual",
        "expected",
        "include_in_eir",
        "principal_consistent",
        "basic_interest",
        "leverage",
        "modified_time_value_pass",
        "prepayment_compensation_reasonable",
        "extension_basic",
        "non_recourse_pass",
        "cli_lookthrough_pass",
        "contingent_feature_basic",
        "de_minimis_or_non_genuine",
        "business_model_change",
        "watchlist",
        "forbearance",
        "unlikely_to_pay",
        "forward_looking_sicr",
        "credit_impaired",
        "low_credit_risk",
        "collective_assessment",
        "sicr_rebuttal",
        "override_approved",
        "separately_recognised",
        "approved",
        "qualitative_substantial",
        "derecognition_approved",
        "risks_rewards_transferred",
        "risks_rewards_retained",
        "control_retained",
        "pass_through_arrangement",
        "partial_transfer",
        "hedging_instrument_eligible",
        "hedged_item_eligible",
        "risk_component_eligible",
        "designation_documented",
        "designated",
        "economic_relationship",
        "credit_risk_dominates",
        "rebalancing_required",
        "discontinue",
        "forecast_transaction_expected",
        "challenger_flag",
        "outcome_flag",
        "nature_dependent_electricity_hedge",
        "variable_nominal_volume",
        "volume_assumptions_documented",
    }
    numbers = {
        "version_no",
        "year",
        "gross_amount",
        "transaction_price",
        "fair_value",
        "transaction_costs",
        "eir",
        "credit_adjusted_eir",
        "reset_rate",
        "undrawn",
        "guarantee_amount",
        "overlay",
        "provision_rate",
        "poci_initial_lifetime_ecl",
        "regulatory_expected_loss",
        "year_fraction",
        "amount",
        "probability",
        "marginal_pd",
        "lifetime_pd_orig",
        "lifetime_pd_current",
        "lgd",
        "ead",
        "ccf",
        "haircut",
        "cost",
        "allocated_amount",
        "recovery_year",
        "fx_rate",
        "total_fv_change",
        "own_credit_change",
        "day_one_difference",
        "hedging_change",
        "hedged_change",
        "hedge_ratio",
        "target_hedge_ratio",
        "option_time_value_change",
        "forward_element_change",
        "basis_spread_change",
        "recycle_amount",
        "basis_adjustment",
        "predicted_value",
        "actual_value",
        "exposure_weight",
        "release_amount",
        "ten_percent_ratio",
        "modified_present_value",
        "continuing_involvement_amount",
        "days_past_due",
        "months_since_cure",
        "months_since_forbearance",
    }
    for col in columns:
        dtype = (
            "boolean"
            if col.startswith("is_") or col in bools
            else "number"
            if col in numbers
            else "string/date"
        )
        dictionary.append(
            [
                col,
                dtype,
                "no"
                if col in {"record_id", "business_key", "version_no", "as_of_date"}
                else "conditional",
                "decimal/EUR as applicable",
                "record_id" if col == "record_id" else "",
                DOMAINS.get(col, ""),
                "business/valid/known as named",
                "canonical source mapping",
                col.replace("_", " "),
            ]
        )
    dictionary.freeze_panes = "A2"
    data = wb.create_sheet(table_name.upper())
    data.append(columns)
    for row in rows:
        data.append([_safe(row.get(col)) for col in columns])
    tab = Table(
        displayName=f"tbl_in_{table_name}"[:250],
        ref=f"A1:{data.cell(1, len(columns)).coordinate}{max(len(rows) + 1, 1)}",
    )
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    data.add_table(tab)
    data.freeze_panes = "A2"
    lookups = wb.create_sheet("_LOOKUPS")
    for idx, (field, domain) in enumerate(DOMAINS.items(), start=1):
        lookups.cell(1, idx, field)
        for ridx, value in enumerate(domain.split("|"), start=2):
            lookups.cell(ridx, idx, value)
    lookups.sheet_state = "hidden"
    changelog = wb.create_sheet("CHANGELOG")
    changelog.append(["version", "date", "change"])
    changelog.append([CONTRACT_VERSION, "2026-09-02", "Initial canonical contract"])
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _normalise_xlsx(path)


def _normalise_xlsx(path: Path) -> None:
    """Make OOXML package timestamps deterministic for reproducible sealing."""
    temp = path.with_suffix(path.suffix + ".normalised")
    with (
        zipfile.ZipFile(path, "r") as source,
        zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target,
    ):
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            info = zipfile.ZipInfo(name, (2026, 9, 2, 10, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-09-02T10:00:00Z\g<2>",
                    content,
                )
            target.writestr(info, content)
    temp.replace(path)


def read_table(path: Path, table_name: str) -> list[dict]:
    wb = load_workbook(path, data_only=False, read_only=False)
    sheet = table_name.upper()
    if sheet not in wb.sheetnames:
        raise ValueError(f"Fehlendes Sheet {sheet} in {path.name}")
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    if not headers or any(x is None for x in headers):
        raise ValueError(f"Ungültiger Header in {path.name}")
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


def _output_table_name(title: str, index: int) -> str:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", title)
    if not safe or safe[0].isdigit():
        safe = "T_" + safe
    return f"tbl_out_{index}_{safe}"[:250]


def write_output(path: Path, sheets: dict[str, list[dict]], run_info: dict | None = None) -> None:
    wb = Workbook()
    fixed_time = datetime(2026, 9, 2, 10, 0, 0)
    wb.properties.created = fixed_time
    wb.properties.modified = fixed_time
    wb.remove(wb.active)
    if run_info is not None:
        sheets = {
            "RUN_INFO": [{"attribute": key, "value": value} for key, value in run_info.items()],
            **sheets,
        }
    if not sheets:
        sheets = {"NOT_APPLICABLE": [{"status": "NOT_APPLICABLE"}]}
    for index, (title, rows) in enumerate(sheets.items(), start=1):
        ws = wb.create_sheet(title[:31])
        headers = list(dict.fromkeys(k for row in rows for k in row)) if rows else ["status"]
        ws.append(headers)
        for row in rows:
            ws.append([_safe(row.get(h)) for h in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        ref = f"A1:{ws.cell(max(ws.max_row, 2), len(headers)).coordinate}"
        if ws.max_row == 1:
            ws.append([None] * len(headers))
        table = Table(displayName=_output_table_name(title, index), ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _normalise_xlsx(path)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, value: dict) -> None:
    path.write_text(
        json.dumps(value, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8"
    )


def manifest_for_files(root: Path, files: list[Path], extra: dict | None = None) -> dict:
    result = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "files": {
            str(p.relative_to(root)): {"sha256": sha256(p), "size": p.stat().st_size}
            for p in sorted(files)
        },
    }
    result.update(extra or {})
    return result
